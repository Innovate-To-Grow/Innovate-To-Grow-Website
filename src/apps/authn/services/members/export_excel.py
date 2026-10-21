"""
Export members to Excel (.xlsx) using openpyxl in write-only mode for performance.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.services.helpers.sheets_safety import safe_sheet_value

# Column definitions: (header_label, width). They match the import template for round-trip compatibility.
_COLUMNS = [
    ("Member UUID", 38),
    ("First Name", 15),
    ("Last Name", 15),
    ("Middle Name", 15),
    ("Title", 20),
    ("Organization", 25),
    ("Active", 10),
    ("Staff", 10),
    ("When Started", 18),
    ("Primary Email", 30),
    ("Primary Verified", 16),
    ("Primary Subscribed", 18),
    ("Secondary Email", 30),
    ("Secondary Verified", 18),
    ("Secondary Subscribed", 20),
    ("Phone Number", 18),
    ("Phone Subscribed", 18),
    ("Phone Verified", 18),
]


def export_members_to_excel(queryset) -> bytes:
    """
    Export a queryset of Member objects to a styled Excel workbook.

    Uses write-only mode and prefetch_related for efficient handling of large datasets.
    Columns match the import template for round-trip compatibility.

    Args:
        queryset: A Django QuerySet of Member objects.

    Returns:
        bytes: The Excel file content.
    """
    queryset = queryset.only(
        "id",
        "first_name",
        "last_name",
        "middle_name",
        "title",
        "organization",
        "is_active",
        "is_staff",
        "date_joined",
    ).prefetch_related("contact_emails", "contact_phones")

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Members")
    _write_header(ws)
    # Data rows — plain values, no per-cell styling for performance
    for member in queryset:
        ws.append(_member_row(member))

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _write_header(ws) -> None:
    """Set the column widths (before any row is appended) and write the styled header row."""
    for col_idx, (_label, width) in enumerate(_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Header styles (matching import template)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    header_cells = []
    for label, _width in _COLUMNS:
        cell = WriteOnlyCell(ws, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        header_cells.append(cell)
    ws.append(header_cells)


def _yes_no(value) -> str:
    return "Yes" if value else "No"


def _member_row(member) -> list:
    contact_emails = member.contact_emails.all()
    primary = next((ce for ce in contact_emails if ce.email_type == "primary"), None)
    secondary = next((ce for ce in contact_emails if ce.email_type == "secondary"), None)
    phone = next(iter(member.contact_phones.all()), None)

    # Member-controlled text (name/title/org/email/phone) is neutralized so
    # a value like ``=HYPERLINK(...)`` is not written as a live formula that
    # executes when a staffer opens the export (CSV/formula injection).
    return [
        str(member.id),
        safe_sheet_value(member.first_name),
        safe_sheet_value(member.last_name),
        safe_sheet_value(member.middle_name or ""),
        safe_sheet_value(member.title or ""),
        safe_sheet_value(member.organization or ""),
        _yes_no(member.is_active),
        _yes_no(member.is_staff),
        member.date_joined.strftime("%Y-%m-%d %H:%M") if member.date_joined else "",
        safe_sheet_value(primary.email_address) if primary else "",
        _yes_no(primary and primary.verified),
        _yes_no(primary and primary.subscribe),
        safe_sheet_value(secondary.email_address) if secondary else "",
        _yes_no(secondary and secondary.verified),
        _yes_no(secondary and secondary.subscribe),
        safe_sheet_value(phone.to_e164()) if phone else "",
        _yes_no(phone and phone.subscribe),
        _yes_no(phone and phone.verified),
    ]
