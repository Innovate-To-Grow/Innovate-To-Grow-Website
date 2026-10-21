from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.event.models import CheckIn, CheckInRecord
from apps.event.tests.helpers import (
    make_event,
    make_member,
    make_registration,
    make_superuser,
    make_ticket,
)


class CheckInAdminTest(TestCase):
    def setUp(self):
        self.admin_user = make_superuser(email="checkin-admin@example.com")
        self.client.login(username="checkin-admin@example.com", password="testpass123")
        self.event = make_event(name="Admin Check-in Event")
        self.check_in = CheckIn.objects.create(event=self.event, name="Main Entrance")

    def test_change_page_shows_live_summary_panel(self):
        response = self.client.get(f"/admin/event/checkin/{self.check_in.pk}/change/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "checkin-summary-config")
        self.assertContains(response, "event/css/checkin_change_summary.css")
        self.assertContains(response, "event/js/checkin_change_summary.js")
        self.assertContains(response, "data-checkin-summary")
        self.assertContains(response, "data-summary-total")
        self.assertContains(response, "data-summary-recent-list")
        self.assertContains(response, "data-summary-lookup-search")
        self.assertContains(response, "data-summary-ticket-filter")
        self.assertContains(response, "data-summary-lookup-list")
        self.assertContains(response, "Registration lookup")
        self.assertContains(response, "All ticket types")
        self.assertContains(response, "Ticket Type")
        self.assertContains(response, "Name, email, organization, title, code")
        self.assertContains(response, "Loading recent scans")
        self.assertContains(response, "pollIntervalMs")
        self.assertContains(response, f"/event/check-in/{self.check_in.pk}/status/")
        self.assertContains(response, "Open Check-in Console")
        self.assertContains(response, "Export Excel")
        self.assertContains(response, reverse("admin:event_checkin_scanner", args=[self.check_in.pk]))
        self.assertContains(response, reverse("admin:event_checkin_export", args=[self.check_in.pk]))
        self.assertNotContains(response, "This station")
        self.assertNotContains(response, "data-summary-station")
        self.assertNotContains(response, "Last 5 at this station")
        self.assertNotContains(response, "VIP Gate")
        self.assertNotContains(response, "Check in records")

    def test_changelist_hides_station_scan_count_column(self):
        ticket = make_ticket(self.event, name="General")
        attendee = make_member(email="list-checked-in@example.com", first_name="Ada", last_name="Lovelace")
        registration = make_registration(attendee, self.event, ticket)
        CheckInRecord.objects.create(check_in=self.check_in, registration=registration)

        response = self.client.get("/admin/event/checkin/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Open Console")
        self.assertNotContains(response, "Scans")

    def test_export_excel_includes_checkin_status_and_member_profile(self):
        ticket = make_ticket(self.event, name="VIP")
        member = make_member(
            email="checkin-export@example.com",
            first_name="Ada",
            last_name="Lovelace",
            organization="Analytical Engines",
            title="Chief Scientist",
        )
        registration = make_registration(
            member,
            self.event,
            ticket,
            attendee_organization="Registration Org",
        )
        CheckInRecord.objects.create(check_in=self.check_in, registration=registration, scanned_by=self.admin_user)

        response = self.client.get(reverse("admin:event_checkin_export", args=[self.check_in.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("checkin_admin-check-in-event_main-entrance_", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            (
                "Event",
                "Check-in",
                "Ticket Type",
                "Ticket Code",
                "Attendee Name",
                "Attendee Email",
                "Attendee Organization",
                "Member Title",
                "Member Organization",
                "Checked In",
                "Check-in Station",
                "Checked-in At",
                "Scanned By",
            ),
        )
        self.assertEqual(
            rows[1][0:11],
            (
                "Admin Check-in Event",
                "Main Entrance",
                "VIP",
                registration.ticket_code,
                "Ada Lovelace",
                "checkin-export@example.com",
                "Registration Org",
                "Chief Scientist",
                "Analytical Engines",
                "Yes",
                "Main Entrance",
            ),
        )

    def test_change_page_live_summary_hides_station_count_when_scans_exist(self):
        ticket = make_ticket(self.event, name="General")
        attendee = make_member(email="checked-in@example.com", first_name="Ada", last_name="Lovelace")
        registration = make_registration(attendee, self.event, ticket)
        CheckInRecord.objects.create(check_in=self.check_in, registration=registration)

        response = self.client.get(f"/admin/event/checkin/{self.check_in.pk}/change/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "data-summary-scanned")
        self.assertNotContains(response, "This station")
        self.assertNotContains(response, "data-summary-station")

    def test_scanner_page_is_available_to_staff(self):
        response = self.client.get(reverse("admin:event_checkin_scanner", args=[self.check_in.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "checkin-console-config")
        self.assertContains(response, "event/css/checkin_console.css")
        self.assertContains(response, "event/js/checkin_console.js")
        self.assertContains(response, "data-sync-status")
        self.assertContains(response, "statusPollIntervalMs")
        self.assertContains(response, "data-camera-message")
        self.assertContains(response, "Request camera access")
        self.assertNotContains(response, "This station")
        self.assertNotContains(response, "data-stat-station")
        self.assertContains(response, f"/event/check-in/{self.check_in.pk}/scan/")
        self.assertContains(response, f"/event/check-in/{self.check_in.pk}/status/")
        self.assertContains(response, f"/event/check-in/{self.check_in.pk}/records/__record_id__/undo/")
        self.assertNotContains(response, "i2g-checkin-state")
        self.assertNotContains(response, "function startCamera")

    def test_scanner_page_rejects_non_staff_user(self):
        user = make_member(email="nonstaff-checkin@example.com")
        self.client.force_login(user)

        response = self.client.get(reverse("admin:event_checkin_scanner", args=[self.check_in.pk]))

        self.assertNotEqual(response.status_code, 200)
