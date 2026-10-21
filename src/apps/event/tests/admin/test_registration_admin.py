import json
from io import BytesIO
from unittest.mock import patch

from django.contrib.admin.helpers import ACTION_CHECKBOX_NAME
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.authn.models import ContactEmail, ContactPhone, Member
from apps.event.admin.registration import EventRegistrationAdmin
from apps.event.models import Event, EventRegistration, Ticket
from apps.event.tests.helpers import (
    make_event,
    make_member,
    make_registration,
    make_superuser,
    make_ticket,
)


@override_settings(ADMIN_REQUIRE_CONFIRMATION=False)
class EventRegistrationAdminTest(TestCase):
    def setUp(self):
        self.admin_user = make_superuser()
        self.client.login(username="admin@example.com", password="testpass123")

    def test_changelist_accessible(self):
        response = self.client.get("/admin/event/eventregistration/")
        self.assertEqual(response.status_code, 200)

    def test_changelist_shows_send_all_ticket_emails_button(self):
        response = self.client.get("/admin/event/eventregistration/")
        self.assertContains(response, "Send All Tickets")
        self.assertContains(response, reverse("admin:event_eventregistration_send_all_ticket_emails"))

    def test_export_column_picker_includes_member_information(self):
        event = make_event()
        ticket = make_ticket(event)
        member = make_member(email="export-picker@example.com")
        registration = make_registration(member, event, ticket)

        response = self.client.post(
            "/admin/event/eventregistration/",
            {
                "action": "export_data",
                ACTION_CHECKBOX_NAME: str(registration.pk),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Member Title")
        self.assertContains(response, "Member Organization")
        self.assertContains(response, "Member Primary Email")
        self.assertContains(response, "Member Phone Numbers")
        self.assertContains(response, "Event Start Date")
        self.assertContains(response, "Event End Date")
        self.assertNotContains(response, 'value="event_date"')

    def test_export_includes_event_start_and_end_dates(self):
        event = make_event(name="Multi-day Export", date="2025-06-15", end_date="2025-06-17")
        ticket = make_ticket(event)
        member = make_member(email="export-dates@example.com")
        registration = make_registration(member, event, ticket)

        response = self.client.post(
            "/admin/event/eventregistration/",
            {
                "action": "export_data",
                ACTION_CHECKBOX_NAME: str(registration.pk),
                "export_confirm": "1",
                "export_format": "json",
                "export_filename": "registration_dates",
                "export_fields": ["event_date", "event_start_date", "event_end_date"],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            json.loads(response.content),
            [
                {
                    "event_date": "2025-06-15",
                    "event_start_date": "2025-06-15",
                    "event_end_date": "2025-06-17",
                }
            ],
        )

    def test_export_accepts_legacy_event_date_selection(self):
        event = make_event(name="Legacy Export", date="2025-06-15", end_date="2025-06-17")
        ticket = make_ticket(event)
        member = make_member(email="legacy-export-date@example.com")
        registration = make_registration(member, event, ticket)

        response = self.client.post(
            "/admin/event/eventregistration/",
            {
                "action": "export_data",
                ACTION_CHECKBOX_NAME: str(registration.pk),
                "export_confirm": "1",
                "export_format": "json",
                "export_fields": ["event_date"],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(json.loads(response.content), [{"event_date": "2025-06-15"}])

    def test_export_old_null_end_date_falls_back_to_start_date(self):
        event = make_event(name="Rolling Deploy Export", date="2025-06-15")
        Event.objects.filter(pk=event.pk).update(end_date=None)
        ticket = make_ticket(event)
        member = make_member(email="export-null-end-date@example.com")
        registration = make_registration(member, event, ticket)

        response = self.client.post(
            "/admin/event/eventregistration/",
            {
                "action": "export_data",
                ACTION_CHECKBOX_NAME: str(registration.pk),
                "export_confirm": "1",
                "export_format": "json",
                "export_filename": "registration_legacy_date",
                "export_fields": ["event_date", "event_start_date", "event_end_date"],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            json.loads(response.content),
            [
                {
                    "event_date": "2025-06-15",
                    "event_start_date": "2025-06-15",
                    "event_end_date": "2025-06-15",
                }
            ],
        )

    def test_export_excel_includes_member_profile_and_contacts(self):
        event = make_event(name="Export Event")
        ticket = make_ticket(event, name="VIP")
        member = make_member(
            email="export-primary@example.com",
            first_name="Ada",
            middle_name="Byron",
            last_name="Lovelace",
            organization="Analytical Engines",
            title="Chief Scientist",
        )
        ContactEmail.objects.create(
            member=member,
            email_address="export-secondary@example.com",
            email_type="secondary",
            verified=True,
        )
        ContactEmail.objects.create(
            member=member,
            email_address="export-other@example.com",
            email_type="other",
            verified=False,
        )
        ContactPhone.objects.create(
            member=member,
            phone_number="2095551212",
            region="1-US",
            verified=True,
        )
        registration = make_registration(
            member,
            event,
            ticket,
            attendee_first_name="Grace",
            attendee_last_name="Hopper",
            attendee_email="grace@example.com",
            attendee_secondary_email="grace.secondary@example.com",
            attendee_phone="2095552323",
            attendee_organization="Navy",
        )

        response = self.client.post(
            "/admin/event/eventregistration/",
            {
                "action": "export_data",
                ACTION_CHECKBOX_NAME: str(registration.pk),
                "export_confirm": "1",
                "export_format": "xlsx",
                "export_filename": "registration_export",
                "export_fields": [
                    "event_name",
                    "ticket_name",
                    "attendee_email",
                    "attendee_organization",
                    "member_full_name",
                    "member_title",
                    "member_organization",
                    "member_primary_email",
                    "member_secondary_emails",
                    "member_other_emails",
                    "member_phone_numbers",
                ],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            (
                "Event Name",
                "Ticket",
                "Attendee Email",
                "Attendee Organization",
                "Member Full Name",
                "Member Title",
                "Member Organization",
                "Member Primary Email",
                "Member Secondary Emails",
                "Member Other Emails",
                "Member Phone Numbers",
            ),
        )
        self.assertEqual(
            rows[1],
            (
                "Export Event",
                "VIP",
                "grace@example.com",
                "Navy",
                "Ada Byron Lovelace",
                "Chief Scientist",
                "Analytical Engines",
                "export-primary@example.com",
                "export-secondary@example.com",
                "export-other@example.com",
                "(209)555-1212",
            ),
        )

    def test_send_all_ticket_emails_confirmation_page(self):
        event = make_event()
        ticket = make_ticket(event)
        member = make_member(email="ticket-confirm@example.com")
        make_registration(member, event, ticket)

        response = self.client.get(reverse("admin:event_eventregistration_send_all_ticket_emails"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Send Ticket Emails to All Registrants")
        self.assertContains(response, "You are about to send ticket emails to 1 registrant")

    @patch("apps.event.services.ticket.mail.send_ticket_email")
    def test_send_all_ticket_emails_posts_all_registrations(self, mock_send):
        event = make_event()
        ticket = make_ticket(event)
        member_one = make_member(email="ticket-one@example.com")
        member_two = make_member(email="ticket-two@example.com")
        registration_one = make_registration(member_one, event, ticket)
        registration_two = make_registration(member_two, event, ticket)

        response = self.client.post(reverse("admin:event_eventregistration_send_all_ticket_emails"))

        self.assertRedirects(response, reverse("admin:event_eventregistration_changelist"))
        self.assertEqual(mock_send.call_count, 2)
        self.assertEqual(
            {call.args[0].pk for call in mock_send.call_args_list},
            {registration_one.pk, registration_two.pk},
        )

    @patch("apps.event.services.ticket.mail.send_ticket_email")
    def test_send_all_ticket_emails_empty_queryset_does_not_send(self, mock_send):
        response = self.client.post(reverse("admin:event_eventregistration_send_all_ticket_emails"))

        self.assertRedirects(response, reverse("admin:event_eventregistration_changelist"))
        mock_send.assert_not_called()

    def test_has_add_permission(self):
        from django.test import RequestFactory

        factory = RequestFactory()
        request = factory.get("/admin/")
        request.user = self.admin_user
        admin_instance = EventRegistrationAdmin(EventRegistration, None)
        self.assertTrue(admin_instance.has_add_permission(request))

    def test_add_form_accessible(self):
        response = self.client.get("/admin/event/eventregistration/add/")
        self.assertEqual(response.status_code, 200)

    @patch("apps.event.services.registration_sheet_sync.schedule_registration_sync")
    def test_admin_add_creates_registration(self, mock_sync):
        member = Member.objects.create_user(password="testpass123")
        ContactEmail.objects.create(
            member=member,
            email_address="reg@e.com",
            email_type="primary",
            verified=True,
        )
        event = make_event()
        ticket = Ticket.objects.create(event=event, name="GA")
        mock_sync.reset_mock()
        response = self.client.post(
            "/admin/event/eventregistration/add/",
            {
                "member": str(member.pk),
                "event": str(event.pk),
                "ticket": str(ticket.pk),
                "attendee_first_name": "",
                "attendee_last_name": "",
                "attendee_email": "",
                "attendee_secondary_email": "",
                "attendee_phone": "",
                "attendee_organization": "",
                "phone_verified": "",
                "question_answers": "[]",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(EventRegistration.objects.filter(member=member, event=event).exists())
        mock_sync.assert_called_once_with(event)

    def test_admin_add_rejects_duplicate(self):
        member = Member.objects.create_user(password="testpass123")
        ContactEmail.objects.create(
            member=member,
            email_address="dup@e.com",
            email_type="primary",
            verified=True,
        )
        event = make_event()
        ticket = Ticket.objects.create(event=event, name="GA")
        EventRegistration.objects.create(member=member, event=event, ticket=ticket)
        response = self.client.post(
            "/admin/event/eventregistration/add/",
            {
                "member": str(member.pk),
                "event": str(event.pk),
                "ticket": str(ticket.pk),
                "attendee_first_name": "",
                "attendee_last_name": "",
                "attendee_email": "",
                "attendee_secondary_email": "",
                "attendee_phone": "",
                "attendee_organization": "",
                "phone_verified": "",
                "question_answers": "[]",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(EventRegistration.objects.filter(member=member, event=event).count(), 1)

    def test_admin_add_rejects_mismatched_ticket(self):
        member = Member.objects.create_user(password="testpass123")
        ContactEmail.objects.create(
            member=member,
            email_address="mis@e.com",
            email_type="primary",
            verified=True,
        )
        event1 = make_event(name="Event 1")
        event2 = make_event(name="Event 2")
        ticket_from_event2 = Ticket.objects.create(event=event2, name="VIP")
        response = self.client.post(
            "/admin/event/eventregistration/add/",
            {
                "member": str(member.pk),
                "event": str(event1.pk),
                "ticket": str(ticket_from_event2.pk),
                "attendee_first_name": "",
                "attendee_last_name": "",
                "attendee_email": "",
                "attendee_secondary_email": "",
                "attendee_phone": "",
                "attendee_organization": "",
                "phone_verified": "",
                "question_answers": "[]",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(EventRegistration.objects.filter(member=member, event=event1).exists())

    def test_has_change_permission(self):
        from django.test import RequestFactory

        factory = RequestFactory()
        request = factory.get("/admin/")
        request.user = self.admin_user
        admin_instance = EventRegistrationAdmin(EventRegistration, None)
        self.assertTrue(admin_instance.has_change_permission(request))

    def test_change_page_keeps_only_ticket_editable_in_ticket_section(self):
        from django.test import RequestFactory

        event = make_event()
        ticket = make_ticket(event)
        member = make_member(email="readonly-ticket-section@example.com")
        registration = make_registration(member, event, ticket)
        factory = RequestFactory()
        request = factory.get("/admin/")
        request.user = self.admin_user
        admin_instance = EventRegistrationAdmin(EventRegistration, None)

        readonly_fields = admin_instance.get_readonly_fields(request, obj=registration)

        self.assertNotIn("ticket", readonly_fields)
        self.assertIn("event", readonly_fields)
        self.assertIn("member", readonly_fields)
        self.assertIn("ticket_code", readonly_fields)
        self.assertIn("attendee_email", readonly_fields)
        self.assertIn("send_ticket_email_action", readonly_fields)

    def test_change_form_limits_ticket_choices_to_registration_event(self):
        event = make_event(name="Ticket Choice Event")
        current_ticket = make_ticket(event, name="Current Ticket")
        new_ticket = make_ticket(event, name="Upgrade Ticket")
        other_event = make_event(name="Other Ticket Event")
        other_ticket = make_ticket(other_event, name="Other Event Ticket")
        member = make_member(email="ticket-choice@example.com")
        registration = make_registration(member, event, current_ticket)

        response = self.client.get(f"/admin/event/eventregistration/{registration.pk}/change/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="ticket"')
        self.assertContains(response, 'name="_send_ticket_email"')
        self.assertContains(response, "Send ticket email now")
        self.assertNotContains(response, 'name="send_ticket_email"')
        self.assertNotContains(response, "If checked, a confirmation email")
        self.assertContains(response, str(new_ticket))
        self.assertNotContains(response, str(other_ticket))

    @patch("apps.event.services.registration_sheet_sync.schedule_registration_sync")
    def test_admin_change_updates_registration_ticket(self, mock_sync):
        event = make_event()
        current_ticket = make_ticket(event, name="General")
        new_ticket = make_ticket(event, name="VIP")
        member = make_member(email="ticket-change@example.com")
        registration = make_registration(member, event, current_ticket)

        mock_sync.reset_mock()
        response = self.client.post(
            f"/admin/event/eventregistration/{registration.pk}/change/",
            {
                "ticket": str(new_ticket.pk),
                "_save": "Save",
            },
        )

        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertEqual(registration.ticket, new_ticket)
        mock_sync.assert_called_once_with(event)

    @patch("apps.event.services.registration_sheet_sync.schedule_registration_sync")
    def test_admin_change_rejects_ticket_from_other_event(self, mock_sync):
        event = make_event(name="Original Event")
        current_ticket = make_ticket(event, name="General")
        other_event = make_event(name="Other Event")
        other_ticket = make_ticket(other_event, name="Other VIP")
        member = make_member(email="ticket-mismatch-change@example.com")
        registration = make_registration(member, event, current_ticket)

        mock_sync.reset_mock()
        response = self.client.post(
            f"/admin/event/eventregistration/{registration.pk}/change/",
            {
                "ticket": str(other_ticket.pk),
                "_save": "Save",
            },
        )

        self.assertEqual(response.status_code, 200)
        registration.refresh_from_db()
        self.assertEqual(registration.ticket, current_ticket)
        mock_sync.assert_not_called()

    @patch("apps.event.services.ticket.mail.send_ticket_email")
    @patch("apps.event.services.registration_sheet_sync.schedule_registration_sync")
    def test_admin_change_send_ticket_button_saves_ticket_and_sends_email(self, mock_sync, mock_send):
        event = make_event()
        current_ticket = make_ticket(event, name="General")
        new_ticket = make_ticket(event, name="VIP")
        member = make_member(email="ticket-send-button@example.com")
        registration = make_registration(member, event, current_ticket)

        mock_sync.reset_mock()
        response = self.client.post(
            f"/admin/event/eventregistration/{registration.pk}/change/",
            {
                "ticket": str(new_ticket.pk),
                "_send_ticket_email": "1",
            },
        )

        self.assertRedirects(response, f"/admin/event/eventregistration/{registration.pk}/change/")
        registration.refresh_from_db()
        self.assertEqual(registration.ticket, new_ticket)
        mock_sync.assert_called_once_with(event)
        mock_send.assert_called_once()
        self.assertEqual(mock_send.call_args.args[0].pk, registration.pk)
        self.assertEqual(mock_send.call_args.args[0].ticket_id, new_ticket.pk)

    def test_has_delete_permission(self):
        from django.test import RequestFactory

        factory = RequestFactory()
        request = factory.get("/admin/")
        request.user = self.admin_user
        admin_instance = EventRegistrationAdmin(EventRegistration, None)
        self.assertTrue(admin_instance.has_delete_permission(request))

    def test_search_by_attendee_fields(self):
        member = Member.objects.create_user(password="testpass123")
        ContactEmail.objects.create(member=member, email_address="u@e.com", email_type="primary", verified=True)
        event = make_event()
        ticket = Ticket.objects.create(event=event, name="GA")
        EventRegistration.objects.create(
            member=member,
            event=event,
            ticket=ticket,
            attendee_first_name="Searchable",
            attendee_last_name="Name",
        )
        response = self.client.get("/admin/event/eventregistration/?q=Searchable")
        self.assertEqual(response.status_code, 200)

    def test_member_info_endpoint(self):
        member = Member.objects.create_user(
            password="testpass123",
            first_name="Ada",
            last_name="Lovelace",
            organization="UCM",
            title="Prof",
        )
        ContactEmail.objects.create(
            member=member,
            email_address="ada@e.com",
            email_type="primary",
            verified=True,
        )
        ContactEmail.objects.create(
            member=member,
            email_address="ada2@e.com",
            email_type="secondary",
            verified=False,
        )
        response = self.client.get(f"/admin/event/eventregistration/member-info/{member.pk}/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["name"], "Ada Lovelace")
        self.assertIn("ada@e.com", data["emails"])
        self.assertIn("ada2@e.com", data["emails"])
        self.assertEqual(data["organization"], "UCM")
        self.assertEqual(data["title"], "Prof")

    def test_event_info_endpoint(self):
        event = make_event(
            name="Info Test",
            allow_secondary_email=True,
            end_date="2025-06-17",
            registration_open=True,
            collect_phone=True,
            verify_phone=True,
            require_phone=True,
            verify_secondary_email=True,
            require_secondary_email=True,
        )
        Ticket.objects.create(event=event, name="GA")
        Ticket.objects.create(event=event, name="VIP")
        response = self.client.get(f"/admin/event/eventregistration/event-info/{event.pk}/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["name"], "Info Test")
        self.assertEqual(data["date"], "2025-06-15")
        self.assertEqual(data["end_date"], "2025-06-17")
        self.assertEqual(data["date_range"], "June 15–17, 2025")
        self.assertEqual(data["location"], "Test Venue")
        self.assertNotIn("is_live", data)
        self.assertTrue(data["registration_open"])
        self.assertTrue(data["collect_phone"])
        self.assertTrue(data["verify_phone"])
        self.assertTrue(data["require_phone"])
        self.assertTrue(data["verify_secondary_email"])
        self.assertTrue(data["require_secondary_email"])
        self.assertEqual(data["total_registrations"], 0)
        ticket_names = [t["name"] for t in data["tickets"]]
        self.assertIn("GA", ticket_names)
        self.assertIn("VIP", ticket_names)
